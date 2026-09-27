import os
import pandas as pd
import networkx as nx
from openai import OpenAI
from dotenv import load_dotenv
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule

load_dotenv()
client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))

print("1. Loading data & building enhanced graph...")
input_path = "data/supply_chain.xlsx"
df_suppliers = pd.read_excel(input_path, sheet_name="Suppliers")
df_shipments = pd.read_excel(input_path, sheet_name="Shipments")
df_incidents = pd.read_excel(input_path, sheet_name="Incidents")

G = nx.Graph()

# Populate Graph Nodes & Edges
for _, row in df_suppliers.iterrows():
    G.add_node(row["SupplierID"], type="Supplier", name=row["SupplierName"], region=row["Region"], risk=row["RiskRating"])

for _, row in df_shipments.iterrows():
    shp_id, sup_id, wh, mat = row["ShipmentID"], row["SupplierID"], row["Warehouse"], row["MaterialCategory"]
    G.add_node(shp_id, type="Shipment", material=mat, delay=row["TransitDelayDays"], incident_flag=row["IncidentReported"])
    G.add_node(wh, type="Warehouse")
    G.add_edge(sup_id, shp_id, relation="SUPPLIES_SHIPMENT")
    G.add_edge(shp_id, wh, relation="DELIVERED_TO")

for _, row in df_incidents.iterrows():
    inc_id, shp_id = row["IncidentID"], row["ShipmentID"]
    G.add_node(inc_id, type="Incident", root_cause=row["RootCause"], financial_impact=row["FinancialImpactUSD"])
    G.add_edge(shp_id, inc_id, relation="HAS_INCIDENT")

# Run NetworkX Graph Analytics (PageRank Centrality)
pagerank_scores = nx.pagerank(G, alpha=0.85)

print("2. Running GraphRAG evaluation with centrality insights...")
analysis_results = []

for _, supplier in df_suppliers.iterrows():
    sup_id = supplier["SupplierID"]
    sup_name = supplier["SupplierName"]
    
    # Get structural PageRank score for the supplier node
    sup_pagerank = pagerank_scores.get(sup_id, 0.0)
    
    # Local neighborhood collection
    subgraph_nodes = set([sup_id])
    for neighbor in G.neighbors(sup_id):
        subgraph_nodes.add(neighbor)
        for sub_neighbor in G.neighbors(neighbor):
            subgraph_nodes.add(sub_neighbor)
            
    context_lines = [
        f"Supplier: {sup_name} (ID: {sup_id}, Initial Risk: {supplier['RiskRating']}, Region: {supplier['Region']})",
        f"Network Centrality (PageRank Score): {sup_pagerank:.4f}"
    ]
    for node in subgraph_nodes:
        if node != sup_id:
            node_data = G.nodes[node]
            context_lines.append(f"- Connected Entity [{node}]: Type={node_data.get('type')}, Attributes={dict(node_data)}")
            
    graph_context_text = "\n".join(context_lines)
    
    prompt = f"""
    You are an expert Enterprise Supply Chain Risk Analyst. Using the localized Knowledge Graph neighborhood and structural PageRank metrics, evaluate this supplier's systemic vulnerability.

    Graph Neighborhood & Analytics:
    {graph_context_text}

    Provide your response in 2 distinct parts separated by a pipe character (|):
    1. Systemic Risk Summary & Network Impact Analysis
    2. Recommended Mitigation Action
    """

    try:
        response = client.chat.completions.create(
            model="gpt-4o-mini",
            messages=[{"role": "user", "content": prompt}],
            temperature=0.2
        )
        ai_output = response.choices[0].message.content.strip()
        risk_summary, mitigation = ai_output.split("|", 1) if "|" in ai_output else (ai_output, "Monitor closely.")
    except Exception as e:
        risk_summary, mitigation = f"Error: {str(e)}", "N/A"

    analysis_results.append({
        "SupplierID": sup_id,
        "SupplierName": sup_name,
        "InitialRisk": supplier["RiskRating"],
        "NetworkCentralityScore": round(sup_pagerank, 4),
        "GraphRAG_Assessment": risk_summary.strip(),
        "Recommended_Mitigation": mitigation.strip()
    })

print("3. Exporting to Professionally Styled Excel Workbook...")
output_path = "data/risk_analysis_output.xlsx"
df_output = pd.DataFrame(analysis_results)

with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
    df_output.to_excel(writer, sheet_name="Risk & Network Analysis", index=False)

# Professional Styling with openpyxl
import openpyxl
wb = openpyxl.load_workbook(output_path)
ws = wb.active

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
thin_border = Border(left=Side(style='thin', color='D9D9D9'),
                     right=Side(style='thin', color='D9D9D9'),
                     top=Side(style='thin', color='D9D9D9'),
                     bottom=Side(style='thin', color='D9D9D9'))

for col in range(1, len(df_output.columns) + 1):
    cell = ws.cell(row=1, column=col)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

for row in range(2, len(df_output) + 2):
    for col in range(1, len(df_output.columns) + 1):
        cell = ws.cell(row=row, column=col)
        cell.border = thin_border
        cell.alignment = Alignment(vertical="top", wrap_text=True)

# Auto-fit column widths
for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = openpyxl.utils.get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 3, 15)

wb.save(output_path)
print(f"Enhanced analytics & styled report saved successfully to: {output_path}")